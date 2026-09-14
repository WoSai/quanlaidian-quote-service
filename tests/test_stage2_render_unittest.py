import io
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.domain.pricing import build_quotation_config
from app.domain.render_pdf import render_pdf
from app.domain.render_xlsx import render_xlsx


CATALOG = Path(__file__).resolve().parent.parent / "references" / "product_catalog_v2.json"


class Stage2RenderTests(unittest.TestCase):
    def test_tier_delivery_fee_and_rights_render(self):
        form = {
            "客户品牌名称": "阶段2渲染测试",
            "餐饮业态": "咖啡",
            "餐饮类型": "轻餐",
            "门店数量": 50,
            "所需功能描述": "连锁供应链、配送中心、VIP售后",
            "门店套餐": "轻餐连锁供应链版",
            "门店增值模块": [],
            "总部模块": ["连锁供应链-配送中心-轻/正餐"],
            "配送中心数量": 1,
            "生产加工中心数量": 0,
            "是否购买VIP售后服务": True,
        }
        config = build_quotation_config(form, {"items": []}, CATALOG, descriptions={})
        wb = load_workbook(io.BytesIO(render_xlsx(config)), data_only=False)

        tier = wb["阶梯报价参考"]
        rows = {tier.cell(row, 3).value: row for row in range(1, tier.max_row + 1)}
        delivery_row = rows["供应链-配送中心交付"]
        self.assertEqual(tier.cell(delivery_row, 5).value, 10000)
        self.assertEqual(tier.cell(delivery_row, 6).value, 10000)
        self.assertEqual(tier.cell(delivery_row, 7).value, 20000)
        self.assertEqual(tier.cell(delivery_row, 8).value, 20000)

        vip_row = rows["Vip售后服务"]
        self.assertEqual(tier.cell(vip_row, 6).value, 5000)
        self.assertEqual(tier.cell(vip_row, 8).value, 9900)

        cover_text = "\n".join(
            str(cell.value)
            for sheet in wb.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        self.assertIn("方式二：外卖包年费", cover_text)
        self.assertIn("收钱吧掌柜端——采购商城", cover_text)

        pdf = render_pdf(config)
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertGreater(len(pdf), 1000)


if __name__ == "__main__":
    unittest.main()
